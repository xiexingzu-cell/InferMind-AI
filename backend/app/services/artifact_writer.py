import csv
import subprocess
import zipfile
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.shared import Inches

from app.config import settings
from app.models.competition_artifact import CompetitionArtifact
from app.models.competition_project import CompetitionProject

MEDIA_TYPES = {
    ".csv": "text/csv",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".txt": "text/plain",
    ".zip": "application/zip",
}

TABLE_EXTENSIONS = {".csv"}
FIGURE_EXTENSIONS = {".jpg", ".jpeg", ".png"}


def _project_dir(project_id: str) -> Path:
    path = Path(settings.competition_storage_path) / project_id
    path.mkdir(parents=True, exist_ok=True)
    return path


def _media_type(path: Path) -> str:
    return MEDIA_TYPES.get(path.suffix.lower(), "application/octet-stream")


def _artifact(project_id: str, stage_id: str, path: Path) -> CompetitionArtifact:
    stored_name = path.relative_to(_project_dir(project_id)).as_posix()
    return CompetitionArtifact(
        id=str(uuid4()),
        project_id=project_id,
        stage_id=stage_id,
        filename=path.name,
        stored_name=stored_name,
        media_type=_media_type(path),
        size_bytes=path.stat().st_size,
    )


def artifacts_from_paths(
    project_id: str, stage_id: str, paths: list[Path]
) -> list[CompetitionArtifact]:
    return [_artifact(project_id, stage_id, path) for path in paths if path.is_file()]


def _add_markdown(document: Document, markdown: str) -> None:
    for line in markdown.splitlines():
        if line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.strip():
            document.add_paragraph(line)


def _add_csv_table(document: Document, path: Path, max_rows: int = 30) -> None:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        return

    document.add_heading(f"表：{path.stem}", level=3)
    rows = rows[:max_rows]
    width = max(1, max(len(row) for row in rows))
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index in range(width):
            table.cell(row_index, column_index).text = row[column_index] if column_index < len(row) else ""


def _add_image(document: Document, path: Path) -> None:
    document.add_heading(f"图：{path.stem}", level=3)
    document.add_picture(str(path), width=Inches(5.8))


def _analysis_assets(project_dir: Path) -> list[Path]:
    runner_dir = project_dir / "runner"
    if not runner_dir.exists():
        return []
    assets: list[Path] = []
    for path in runner_dir.rglob("*"):
        if path.is_file() and path.suffix.lower() in TABLE_EXTENSIONS | FIGURE_EXTENSIONS | {".md", ".txt"}:
            assets.append(path)
    return sorted(assets, key=lambda item: item.as_posix())


def _write_results_package(project_dir: Path) -> Path | None:
    assets = _analysis_assets(project_dir)
    if not assets:
        return None
    package_path = project_dir / "analysis-results.zip"
    with zipfile.ZipFile(package_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in assets:
            archive.write(path, arcname=path.relative_to(project_dir).as_posix())
    return package_path


def write_stage_artifacts(
    project: CompetitionProject, stage_id: str, output: str
) -> list[CompetitionArtifact]:
    project_dir = _project_dir(project.id)
    markdown_path = project_dir / f"{stage_id}.md"
    markdown_path.write_text(output, encoding="utf-8")
    artifact_paths = [markdown_path]

    if stage_id != "paper":
        return artifacts_from_paths(project.id, stage_id, artifact_paths)

    document = Document()
    document.add_heading(project.title, level=0)
    _add_markdown(document, output)

    assets = _analysis_assets(project_dir)
    tables = [path for path in assets if path.suffix.lower() in TABLE_EXTENSIONS]
    figures = [path for path in assets if path.suffix.lower() in FIGURE_EXTENSIONS]
    notes = [path for path in assets if path.suffix.lower() in {".md", ".txt"}]

    if tables or figures or notes:
        document.add_page_break()
        document.add_heading("数据分析产物", level=1)
        for note in notes:
            document.add_heading(f"结果摘要：{note.stem}", level=2)
            _add_markdown(document, note.read_text(encoding="utf-8", errors="replace"))
        for table_path in tables:
            _add_csv_table(document, table_path)
        for figure_path in figures:
            _add_image(document, figure_path)

    docx_path = project_dir / "competition-report.docx"
    document.save(docx_path)
    artifact_paths.append(docx_path)

    package_path = _write_results_package(project_dir)
    if package_path is not None:
        artifact_paths.append(package_path)

    try:
        subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(project_dir),
                str(docx_path),
            ],
            check=True,
            timeout=60,
            capture_output=True,
        )
    except (FileNotFoundError, subprocess.SubprocessError):
        return artifacts_from_paths(project.id, stage_id, artifact_paths)

    pdf_path = project_dir / "competition-report.pdf"
    if pdf_path.exists():
        artifact_paths.append(pdf_path)
    return artifacts_from_paths(project.id, stage_id, artifact_paths)
